"""
Excel Export Service using openpyxl
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import logging

logger = logging.getLogger(__name__)


class ExcelExportService:
    """Service for exporting data to Excel"""

    def __init__(self):
        self.header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF", size=12)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def export_customers(self, customers, output_path: str):
        """Export customers to Excel"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Customers"

            # Headers
            headers = [
                "ID", "Full Name", "Email", "Phone", "Status", "Risk Profile",
                "Portfolio Value", "Created Date"
            ]
            ws.append(headers)

            # Style headers
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = self.header_font
                cell.fill = self.header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self.border

            # Add data
            for customer in customers:
                ws.append([
                    customer.id,
                    customer.full_name,
                    customer.email,
                    customer.phone,
                    customer.status.value if customer.status else "",
                    customer.risk_profile.value if customer.risk_profile else "",
                    customer.current_portfolio_value or 0,
                    customer.created_at.strftime("%Y-%m-%d") if customer.created_at else ""
                ])

            # Auto-size columns
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            wb.save(output_path)
            logger.info(f"Customers exported to Excel: {output_path}")
            return output_path

        except Exception as e:
            logger.error(f"Error exporting customers to Excel: {e}")
            raise

    def export_commissions(self, commissions, output_path: str):
        """Export commissions to Excel"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Commissions"

            # Headers
            headers = [
                "ID", "Customer Name", "Type", "Amount", "Percentage",
                "Earned Date", "Payment Date", "Status", "Notes"
            ]
            ws.append(headers)

            # Style headers
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = self.header_font
                cell.fill = self.header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self.border

            # Add data
            for commission in commissions:
                ws.append([
                    commission.id,
                    commission.customer.full_name if commission.customer else "",
                    commission.commission_type,
                    commission.amount,
                    commission.percentage or "",
                    commission.earned_date.strftime("%Y-%m-%d") if commission.earned_date else "",
                    commission.payment_date.strftime("%Y-%m-%d") if commission.payment_date else "",
                    "PAID" if commission.is_paid else "PENDING",
                    commission.notes or ""
                ])

            # Auto-size columns
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            # Add totals row
            total_row = ws.max_row + 2
            ws.cell(row=total_row, column=3, value="TOTAL:")
            ws.cell(row=total_row, column=3).font = Font(bold=True)
            ws.cell(row=total_row, column=4, value=f"=SUM(D2:D{ws.max_row-1})")
            ws.cell(row=total_row, column=4).font = Font(bold=True)

            wb.save(output_path)
            logger.info(f"Commissions exported to Excel: {output_path}")
            return output_path

        except Exception as e:
            logger.error(f"Error exporting commissions to Excel: {e}")
            raise

    def export_invoices(self, invoices, output_path: str):
        """Export invoices to Excel"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Invoices"

            # Headers
            headers = [
                "Invoice Number", "Customer Name", "Invoice Date", "Due Date",
                "Subtotal", "Tax", "Total Amount", "Paid Amount", "Balance",
                "Status", "Payment Date", "Payment Method"
            ]
            ws.append(headers)

            # Style headers
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = self.header_font
                cell.fill = self.header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self.border

            # Add data
            for invoice in invoices:
                balance = invoice.total_amount - invoice.paid_amount
                ws.append([
                    invoice.invoice_number,
                    invoice.customer.full_name if invoice.customer else "",
                    invoice.invoice_date.strftime("%Y-%m-%d") if invoice.invoice_date else "",
                    invoice.due_date.strftime("%Y-%m-%d") if invoice.due_date else "",
                    invoice.subtotal,
                    invoice.tax_amount,
                    invoice.total_amount,
                    invoice.paid_amount,
                    balance,
                    invoice.status.value if invoice.status else "",
                    invoice.payment_date.strftime("%Y-%m-%d") if invoice.payment_date else "",
                    invoice.payment_method or ""
                ])

            # Auto-size columns
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            # Add totals row
            total_row = ws.max_row + 2
            ws.cell(row=total_row, column=4, value="TOTAL:")
            ws.cell(row=total_row, column=4).font = Font(bold=True)
            for col in [5, 6, 7, 8, 9]:  # Subtotal, Tax, Total, Paid, Balance
                ws.cell(row=total_row, column=col, value=f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{ws.max_row-1})")
                ws.cell(row=total_row, column=col).font = Font(bold=True)

            wb.save(output_path)
            logger.info(f"Invoices exported to Excel: {output_path}")
            return output_path

        except Exception as e:
            logger.error(f"Error exporting invoices to Excel: {e}")
            raise

    def export_investments(self, investments, output_path: str):
        """Export investments to Excel"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Investments"

            # Headers
            headers = [
                "ID", "Customer Name", "Investment Name", "Category",
                "Invested Amount", "Current Value", "Returns", "Return %",
                "Investment Date", "Description"
            ]
            ws.append(headers)

            # Style headers
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = self.header_font
                cell.fill = self.header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self.border

            # Add data
            for investment in investments:
                current_value = investment.current_value or investment.invested_amount
                returns = current_value - investment.invested_amount
                return_pct = (returns / investment.invested_amount * 100) if investment.invested_amount > 0 else 0

                ws.append([
                    investment.id,
                    investment.customer.full_name if investment.customer else "",
                    investment.investment_name,
                    investment.investment_type.value if investment.investment_type else "",
                    investment.invested_amount,
                    current_value,
                    returns,
                    f"{return_pct:.2f}%",
                    investment.investment_date.strftime("%Y-%m-%d") if investment.investment_date else "",
                    investment.description or ""
                ])

            # Auto-size columns
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            # Add totals row
            total_row = ws.max_row + 2
            ws.cell(row=total_row, column=4, value="TOTAL:")
            ws.cell(row=total_row, column=4).font = Font(bold=True)
            for col in [5, 6, 7]:  # Invested, Current, Returns
                ws.cell(row=total_row, column=col, value=f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{ws.max_row-1})")
                ws.cell(row=total_row, column=col).font = Font(bold=True)

            wb.save(output_path)
            logger.info(f"Investments exported to Excel: {output_path}")
            return output_path

        except Exception as e:
            logger.error(f"Error exporting investments to Excel: {e}")
            raise

    def export_prospects(self, prospects, output_path: str):
        """Export prospects to Excel"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Prospects"

            # Headers
            headers = [
                "ID", "Full Name", "Email", "Phone", "Company", "Designation",
                "Status", "Priority", "Source", "Estimated Value", "Assigned To",
                "Next Follow-up", "Last Contact", "Contact Attempts", "Tags",
                "Notes", "Created Date"
            ]
            ws.append(headers)

            # Style headers
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = self.header_font
                cell.fill = self.header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = self.border

            # Add data
            for prospect in prospects:
                # Format tags as comma-separated string
                tags_str = ", ".join(prospect.tags) if prospect.tags else ""

                # Get assigned user name
                assigned_to = prospect.assigned_to.full_name if prospect.assigned_to else ""

                ws.append([
                    prospect.id,
                    prospect.full_name,
                    prospect.email or "",
                    prospect.phone or "",
                    prospect.company or "",
                    prospect.designation or "",
                    prospect.status.value if prospect.status else "",
                    prospect.priority.value if prospect.priority else "",
                    prospect.source.value if prospect.source else "",
                    prospect.estimated_portfolio_value or "",
                    assigned_to,
                    prospect.next_follow_up_date.strftime("%Y-%m-%d %H:%M") if prospect.next_follow_up_date else "",
                    prospect.last_contact_date.strftime("%Y-%m-%d %H:%M") if prospect.last_contact_date else "",
                    prospect.contact_attempts or 0,
                    tags_str,
                    prospect.notes or "",
                    prospect.created_at.strftime("%Y-%m-%d") if prospect.created_at else ""
                ])

            # Auto-size columns
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            # Add summary statistics at the bottom
            summary_row = ws.max_row + 3
            ws.cell(row=summary_row, column=1, value="Summary Statistics:")
            ws.cell(row=summary_row, column=1).font = Font(bold=True, size=14)

            summary_row += 1
            ws.cell(row=summary_row, column=1, value="Total Prospects:")
            ws.cell(row=summary_row, column=2, value=len(prospects))
            ws.cell(row=summary_row, column=2).font = Font(bold=True)

            # Count by status
            from collections import Counter
            status_counts = Counter([p.status.value for p in prospects if p.status])
            summary_row += 2
            ws.cell(row=summary_row, column=1, value="By Status:")
            ws.cell(row=summary_row, column=1).font = Font(bold=True)
            for status, count in status_counts.items():
                summary_row += 1
                ws.cell(row=summary_row, column=2, value=f"{status}:")
                ws.cell(row=summary_row, column=3, value=count)

            wb.save(output_path)
            logger.info(f"Prospects exported to Excel: {output_path}")
            return output_path

        except Exception as e:
            logger.error(f"Error exporting prospects to Excel: {e}")
            raise


# Global instance
excel_export_service = ExcelExportService()
